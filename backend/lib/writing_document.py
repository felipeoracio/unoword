"""Bounded Tiptap JSON validation and safe portable document representations."""

import json
import re
from html import escape
from io import BytesIO
from urllib.parse import urlsplit

from fastapi import HTTPException

EMPTY = {"type": "doc", "content": [{"type": "paragraph"}]}
NODES = {
    "doc",
    "paragraph",
    "heading",
    "text",
    "bulletList",
    "orderedList",
    "listItem",
    "blockquote",
    "horizontalRule",
    "hardBreak",
}
MARKS = {"bold", "italic", "underline", "strike", "link"}


def safe_link(value):
    try:
        url = urlsplit(str(value))
        return (
            str(value)
            if url.scheme in {"https", "http", "mailto"}
            and not re.search(r"[\x00-\x20]", str(value))
            else None
        )
    except ValueError:
        return None


def clean_document(value):
    count = 0

    def visit(node, depth=0):
        nonlocal count
        count += 1
        if (
            depth > 30
            or count > 15000
            or not isinstance(node, dict)
            or node.get("type") not in NODES
        ):
            raise HTTPException(422, "unsupported_document_structure")
        kind = node["type"]
        out = {"type": kind}
        if kind == "text":
            out["text"] = str(node.get("text", ""))
            marks = []
            for mark in node.get("marks", [])[:10]:
                if not isinstance(mark, dict) or mark.get("type") not in MARKS:
                    continue
                if mark["type"] == "link":
                    href = safe_link(mark.get("attrs", {}).get("href", ""))
                    if href:
                        marks.append(
                            {
                                "type": "link",
                                "attrs": {
                                    "href": href,
                                    "target": "_blank",
                                    "rel": "noopener noreferrer",
                                },
                            }
                        )
                else:
                    marks.append({"type": mark["type"]})
            if marks:
                out["marks"] = marks
        else:
            attrs = node.get("attrs") or {}
            allowed = {}
            if kind == "heading":
                allowed["level"] = max(1, min(3, int(attrs.get("level", 1))))
            if kind in {"heading", "paragraph"} and attrs.get("textAlign") in {
                "left",
                "center",
                "right",
                "justify",
            }:
                allowed["textAlign"] = attrs["textAlign"]
            if kind == "orderedList":
                allowed["start"] = max(1, min(10000, int(attrs.get("start", 1))))
            if allowed:
                out["attrs"] = allowed
            if "content" in node:
                if not isinstance(node["content"], list):
                    raise HTTPException(422, "invalid_document")
                out["content"] = [visit(child, depth + 1) for child in node["content"]]
                blocks = {
                    "paragraph",
                    "heading",
                    "bulletList",
                    "orderedList",
                    "blockquote",
                    "horizontalRule",
                }
                allowed_children = {
                    "doc": blocks,
                    "blockquote": blocks,
                    "listItem": blocks,
                    "paragraph": {"text", "hardBreak"},
                    "heading": {"text", "hardBreak"},
                    "bulletList": {"listItem"},
                    "orderedList": {"listItem"},
                }
                if any(
                    child["type"] not in allowed_children.get(kind, set())
                    for child in out["content"]
                ):
                    raise HTTPException(422, "invalid_document_structure")
        return out

    try:
        if not isinstance(value, dict) or value.get("type") != "doc":
            raise HTTPException(422, "invalid_document")
        if len(json.dumps(value).encode()) > 1048576:
            raise HTTPException(413, "document_too_large")
        cleaned = visit(value)
        if len(plain_text(cleaned)) > 200000:
            raise HTTPException(413, "document_too_large")
        return cleaned
    except (TypeError, ValueError, AttributeError, RecursionError) as error:
        raise HTTPException(422, "invalid_document") from error


def plain_text(node):
    if node.get("type") == "text":
        return node.get("text", "")
    if node.get("type") == "hardBreak":
        return "\n"
    sep = (
        "\n"
        if node.get("type")
        in {"doc", "bulletList", "orderedList", "listItem", "blockquote"}
        else ""
    )
    return sep.join(plain_text(child) for child in node.get("content", []))


def markdown(node, depth=0):
    kind = node["type"]
    if kind == "text":
        text = re.sub(r"([\\`*_\[\]<>])", r"\\\1", node.get("text", ""))
        for mark in node.get("marks", []):
            name = mark["type"]
            if name == "link":
                text = f'[{text}]({mark["attrs"]["href"].replace(")","%29")})'
            elif name in {"bold", "italic", "strike"}:
                wrapper = {"bold": "**", "italic": "*", "strike": "~~"}[name]
                text = wrapper + text + wrapper
        return text
    if kind == "hardBreak":
        return "  \n"
    if kind == "horizontalRule":
        return "\n---\n"
    children = node.get("content", [])
    if kind in {"bulletList", "orderedList"}:
        start = node.get("attrs", {}).get("start", 1)
        return (
            "\n".join(
                ("  " * depth)
                + (f"{start+i}. " if kind == "orderedList" else "- ")
                + markdown(child, depth + 1).strip()
                for i, child in enumerate(children)
            )
            + "\n"
        )
    text = ("\n\n" if kind in {"doc", "listItem", "blockquote"} else "").join(
        markdown(child, depth) for child in children
    )
    if kind == "heading":
        return "#" * node.get("attrs", {}).get("level", 1) + " " + text
    if kind == "blockquote":
        return "\n".join("> " + line for line in text.splitlines())
    return text


def to_docx(content, title):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE

    document = Document()

    def block(node, list_style=None):
        kind = node["type"]
        if kind in {"doc", "bulletList", "orderedList", "blockquote", "listItem"}:
            style = (
                "List Bullet"
                if kind == "bulletList"
                else (
                    "List Number"
                    if kind == "orderedList"
                    else "Quote" if kind == "blockquote" else list_style
                )
            )
            for child in node.get("content", []):
                block(child, style)
            return
        paragraph = document.add_paragraph(style=list_style)
        if kind == "heading":
            paragraph.style = f'Heading {node.get("attrs",{}).get("level",1)}'
        alignment = node.get("attrs", {}).get("textAlign")
        if alignment:
            paragraph.alignment = {
                "left": WD_ALIGN_PARAGRAPH.LEFT,
                "center": WD_ALIGN_PARAGRAPH.CENTER,
                "right": WD_ALIGN_PARAGRAPH.RIGHT,
                "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
            }[alignment]
        if kind == "horizontalRule":
            paragraph.add_run("────────────────")
            return
        for child in node.get("content", []):
            run = paragraph.add_run(
                "\n" if child["type"] == "hardBreak" else plain_text(child)
            )
            marks = {mark["type"] for mark in child.get("marks", [])}
            run.bold = "bold" in marks
            run.italic = "italic" in marks
            run.underline = "underline" in marks
            run.font.strike = "strike" in marks
            link = next(
                (mark for mark in child.get("marks", []) if mark["type"] == "link"),
                None,
            )
            if link:
                relationship = paragraph.part.relate_to(
                    link["attrs"]["href"], RELATIONSHIP_TYPE.HYPERLINK, is_external=True
                )
                hyperlink = OxmlElement("w:hyperlink")
                hyperlink.set(qn("r:id"), relationship)
                paragraph._p.remove(run._r)
                hyperlink.append(run._r)
                paragraph._p.append(hyperlink)

    block(content)
    document.core_properties.title = title
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def to_pdf(content, title):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
    from reportlab.lib.styles import getSampleStyleSheet

    styles = getSampleStyleSheet()
    story = []

    def inline(node):
        if node["type"] == "hardBreak":
            return "<br/>"
        text = escape(node.get("text", "")).replace("\n", "<br/>")
        for mark in node.get("marks", []):
            wrapper = {
                "bold": "b",
                "italic": "i",
                "underline": "u",
                "strike": "strike",
            }.get(mark["type"])
            if wrapper:
                text = f"<{wrapper}>{text}</{wrapper}>"
            elif mark["type"] == "link":
                text = f'<link href="{escape(mark["attrs"]["href"],quote=True)}" color="blue">{text}</link>'
        return text

    def block(node, prefix=None, quoted=False):
        kind = node["type"]
        if kind == "horizontalRule":
            story.extend([HRFlowable(width="100%", color="#b8c5d9"), Spacer(1, 12)])
            return
        if kind in {"doc", "bulletList", "orderedList", "listItem", "blockquote"}:
            for index, child in enumerate(node.get("content", [])):
                marker = (
                    f'{node.get("attrs",{}).get("start",1)+index}.'
                    if kind == "orderedList"
                    else "•" if kind == "bulletList" else prefix
                )
                block(child, marker, quoted or kind == "blockquote")
            return
        style = (
            styles[f'Heading{node.get("attrs",{}).get("level",1)}']
            if kind == "heading"
            else styles["BodyText"]
        )
        # Escape all user text before ReportLab's small markup parser.
        style = style.clone(
            "writing",
            alignment={"left": 0, "center": 1, "right": 2, "justify": 4}.get(
                node.get("attrs", {}).get("textAlign"), 0
            ),
            leftIndent=18 if prefix or quoted else 0,
        )
        text = "".join(inline(child) for child in node.get("content", []))
        story.extend(
            [
                Paragraph(text or " ", style, bulletText=prefix),
                Spacer(1, 8),
            ]
        )

    block(content)
    output = BytesIO()
    SimpleDocTemplate(output, title=title, author="UnoWord").build(story)
    return output.getvalue()


def import_docx(data):
    from docx import Document
    from lib.document_processing import _validate_docx_archive

    _validate_docx_archive(data)
    document = Document(BytesIO(data))
    blocks = []
    for paragraph in document.paragraphs:
        inline = []
        from docx.text.hyperlink import Hyperlink

        runs = []
        for part in paragraph.iter_inner_content():
            if isinstance(part, Hyperlink):
                runs.extend((run, safe_link(part.url)) for run in part.runs)
            else:
                runs.append((part, None))
        for run, href in runs:
            if not run.text:
                continue
            marks = [
                {"type": name}
                for name, value in [
                    ("bold", run.bold),
                    ("italic", run.italic),
                    ("underline", run.underline),
                    ("strike", run.font.strike),
                ]
                if value
            ]
            if href:
                marks.append({"type": "link", "attrs": {"href": href}})
            inline.append(
                {
                    "type": "text",
                    "text": run.text,
                    **({"marks": marks} if marks else {}),
                }
            )
        node = {"type": "paragraph", "content": inline}
        name = paragraph.style.name if paragraph.style else ""
        if name.startswith("Heading ") and name[-1:] in {"1", "2", "3"}:
            node.update(type="heading", attrs={"level": int(name[-1])})
        if paragraph.alignment in {0, 1, 2, 3}:
            node.setdefault("attrs", {})["textAlign"] = {
                0: "left",
                1: "center",
                2: "right",
                3: "justify",
            }[paragraph.alignment]
        if name in {"List Bullet", "List Number"}:
            list_kind = "bulletList" if name == "List Bullet" else "orderedList"
            if not blocks or blocks[-1]["type"] != list_kind:
                blocks.append({"type": list_kind, "content": []})
            blocks[-1]["content"].append({"type": "listItem", "content": [node]})
        else:
            blocks.append(node)
    for table in document.tables:
        for row in table.rows:
            blocks.append(
                {
                    "type": "paragraph",
                    "content": [
                        {
                            "type": "text",
                            "text": " | ".join(cell.text for cell in row.cells),
                        }
                    ],
                }
            )
    return clean_document({"type": "doc", "content": blocks or EMPTY["content"]})
